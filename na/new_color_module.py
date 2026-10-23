# -*- coding: utf-8 -*-
"""
新色跟进判定内核 v2 (2026-09-18 与用户锁定)

母体三态 (按 store×SPU 互斥):
  ① 该店铺该 SPU 完全无 Listing 记录            → 跳过, 无新色
  ② 该店铺该 SPU 仅有"已删除"SKU(无在售/停售)  → 跳过, 无新色
  ③ 该店铺该 SPU 有 ≥1 个非已删除 SKU          → 进入新色判定

判定内核 (仅情形③):
  Amazon 已上架基准 = 在售 ∪ 停售  (停售计入, 已删除排除)
  Walmart 已上架基准 = PUBLISHED ∪ UNPUBLISHED (两者都算有链接, 对称套用母体三态)
  待上架新色 = (在仓[库存清单] + 在途[动态库存表] 出现的颜色) − 基准
            → 已删除 / 非PUBLISHED-UNPUBLISHED 颜色不算上架, 有货即列为待上架新色(含重新上架)

店铺口径: 分店铺判定 (A店已上/ B店未上 → B店仍算新色)
对比粒度: 颜色级 (复用 extract_spu_from_sku + parse_sku_color_suffix, 去掉 A/B/C 包裹后缀)

v3 (2026-09-21 与用户确认):
  1) 套件链路固定为「上架SKU(平台MSKU) → SKU平台信息表 按 店铺+14K货号 查子件(成员货品)」
     例: HulalahomeDeco-US / BCRB1854-BLK-AB-S2-DEM1
         → US-Amazon-HLLdeco-KH 子件 BCRB1854-BLK-B + BCRB1854-BLK-A
         → 该店 BCRB1854-BLK-B / BCRB1854-BLK-A 视为已上架
     Walmart 上架表店铺直接对应套件表 US-Walmart
  2) 展示维度: 新色(未上架) 按「库存SKU」维度一行;
     已上架 按「上架SKU + 库存SKU」维度一对一行 (同一库存SKU挂多个平台SKU时展开多行)
  3) 店铺显示名: HLLdeco→DECO / Home Nest→HN / Urban→Urban / XYT→XYT / JYT→JYT / US-Walmart→Walmart

v4 (2026-09-21 颜色档案 archive-first, 与用户确认):
  颜色判定只认「系统存档」, 不再做任何自动猜测:
    ① 颜色档案 = 内置种子(color_dict_seed.py, 2026-09-21 用户核对表281条确认)
                 + 云端/本地 color_dict.json (网页端人工确认, 优先级更高)
    ② status=颜色   → 色码映射为规范色名(前后缀/规格前缀合并到本体色), 正常参与基准对比
    ③ status=非颜色 → 色码剔除, 不算新色 (明细行上架状态='非颜色', 灰色显示)
    ④ 档案中不存在的新颜色码 → 保留原码照常判定, 标记 颜色判定='待确认',
       汇总到 pending_colors, 网页端「颜色待确认」面板人工确认后存档,
       下次处理生效 → 除非人工调整, 判定始终沿用系统历史数据

v5 (2026-09-21 与用户确认):
  1) SPU 主行 = 一个 SPU × 一个店铺 (同SPU多店铺有上架 → 多条记录),
     新色数/新色列表/待上架SKU数 均按该店铺口径; 折叠展开只显示该店铺该SPU的SKU
  2) 店铺运营规则 (运营组/运营):
       Walmart        → 综合组-WLM · Chloris            (固定)
       XYT / JYT      → 品牌组 · MIA黄                   (固定)
       Urban          → 该SPU款式表「综合组-AMZ」行的运营 (运营组显示 综合组-AMZ)
       DECO / HN      → 该SPU款式表「非综合组-AMZ」行的运营 (软包组/板式组/品牌组…)
     占位: 「项目组」运营显示 '/'
     兜底① 该SPU款式表原值; 兜底② 本体SPU款式表完全无记录时, 用「上架SKU解析出的SPU」
           再查一遍 (如 BDJS0893-KB 的 Urban 上架SKU 是 BDJH0885-KB-xxx → 按 BDJH0885 取)

v6 (2026-10-06 与用户确认: 生命周期 / 淘汰 收敛):
  淘汰口径: 生命周期含「淘汰」 或 等于「下市期」 → 已淘汰 (死货/小黑屋/空值 不淘汰)
  生命周期取值优先级(按SKU): 动态库存表[在途] → 库存清单[在库, 有列才用] → 套件表[美国生命周期状态]
  1) 规则1(淘汰色): 在库/在途池里 生命周期已淘汰的 SKU → 标 上架状态='已淘汰', 不列为待上架新色
  2) 规则2(套件全员非淘汰): 库存SKU 若是某上架SKU(平台SKU)的套件成员, 其所属套件组(全部库存子SKU)
     必须全部非淘汰, 才算需要上架的新色; 否则标 上架状态='套件淘汰' (如 Z6LRSFBR0692-BGE-DEM1
     的成员 Z6SFXW0619-BGE-A / ZBRQX0201-IDG 任一淘汰 → 该新色不计入)
  3) 规则4(母体资格收紧): 已上架基准只计入「非删除 且 非淘汰」的颜色; 某 SPU×店铺 仅已删除 /
     仅淘汰 / 无有效在售 → baseline 为空 → 整行不进新色跟进 (自然也无需上架新色)
  4) 上架SKU 增加「后台状态」字段(原始 Listing/ItemReport 状态: 在售/停售/PUBLISHED/UNPUBLISHED…),
     供网页端「上架SKU 汇总」表格展示
"""
import os
import glob
from collections import defaultdict

import pandas as pd

# Listing 店铺名 → 套件文件店铺名 (+显示名)
AMAZON_STORE_MAP = {
    'HulalahomeDeco-US': ('US-Amazon-HLLdeco-KH', 'DECO'),
    'HOME NEST-US': ('US-Amazon-Home Nest-KH', 'HN'),
    'Urban Living-US': ('US-Amazon-Urban-KH', 'Urban'),
    'XYT HOME-US': ('US-Amazon-XYT-KH', 'XYT'),
    'JYT HOME STORE-US': ('US-Amazon-JYT-KH', 'JYT'),
}
WALMART_STORE = ('US-Walmart', 'Walmart')

# 店铺 key -> 显示名
STORE_LABELS = {v[0]: v[1] for v in AMAZON_STORE_MAP.values()}
STORE_LABELS[WALMART_STORE[0]] = WALMART_STORE[1]

# 店铺展示顺序
STORE_ORDER = ['DECO', 'HN', 'Urban', 'XYT', 'JYT', 'Walmart']

# 内置非颜色(配件)色码: SKU 色段为这些值时判定为非颜色, 不参与新色判定 (如 26SFXW0526-Hardware)
_BUILTIN_NON_COLORS = {'HARDWARE'}

# ---- 店铺运营规则 (2026-09-21 用户确认) ----
# 固定映射 (不查款式表)
STORE_FIXED_OP = {
    'Walmart': ('综合组-WLM', 'Chloris'),
    'XYT': ('品牌组', 'MIA黄'),
    'JYT': ('品牌组', 'MIA黄'),
}
# 非综合组-AMZ 中不代表真实运营组的取值 (取"非综合组-AMZ运营"时排除)
FAKE_GROUP = {'不上架', '已下架', '套组不上', '项目组'}
FAKE_OPERATOR = {'', '/', '-', '无', '待定', '不上架', '已下架', '套组不上'}

# ---- 颜色缩写 ↔ 全称 归一表 (2026-09-21 用户要求: 缩写与全称必须匹配上) ----
# 只映射到「档案里已存在的规范色名」, 不会自行发明新色名;
# 档案本身的映射优先级最高, 这里只兜住档案没写到的写法 (如库存SKU用全称 BEM0908-brown)。
STANDARD_COLOR_VARIANTS = {
    'BLACK': ['BLK', 'BLA', 'BLACK'],
    'WHITE': ['WHT', 'WHI', 'WHITE'],
    'BROWN': ['BRN', 'BRW', 'BROWN'],
    'DARK BROWN': ['DBN', 'DKBRN', 'DARK BROWN'],
    'REDDISH BROWN': ['RBR', 'REDDISH BROWN'],
    'LIGHT BROWN': ['LBR', 'LTBRN', 'LIGHT BROWN'],
    'GREY': ['GRY', 'GRAY', 'GREY'],
    'LIGHT GREY': ['LGR', 'LTGRY', 'LIGHT GREY', 'LIGHT GRAY'],
    'DARK GREY': ['DGR', 'DKYGRY', 'DARK GREY', 'DARK GRAY'],
    'BLUE GREY': ['BLUGY', 'BLUE GREY', 'BLUE GRAY'],
    'ALMOND GRAY': ['ALMOND GRAY', 'ALMGY'],
    'MOONLIGHT GRAY': ['MOONLIGHT GRAY', 'MLGY'],
    'CHARCOAL': ['CHC', 'CHA', 'CHR', 'CHARCOAL'],
    'IVORY': ['IVY', 'IVR', 'IVORY'],
    'IVORY TAUPE': ['IVORY TAUPE', 'IVTP'],
    'BEIGE': ['BGE', 'BEI', 'BEIGE'],
    'CREAM': ['CRM', 'CRE', 'CREAM'],
    'CAMEL': ['CML', 'CAMEL'],
    'CARAMEL': ['CAR', 'CML', 'CARAMEL'],
    'SILVER': ['SLV', 'SIL', 'SILVER'],
    'GOLDEN': ['GLD', 'GOD', 'GOLD', 'GOLDEN'],
    'BLUE': ['BLU', 'BLUE'],
    'NAVY': ['NVY', 'NAV', 'NAVY'],
    'GREEN': ['GRN', 'GREEN'],
    'DARK GREEN': ['DGN', 'DKG', 'DARK GREEN'],
    'MOSS': ['MOS', 'MOSS'],
    'SAGE': ['SGE', 'SAG', 'SAGE'],
    'RED': ['RD', 'RED'],
    'BURGUNDY': ['BUR', 'BURG', 'BURGUNDY'],
    'CORAL': ['CRL', 'COR', 'CORAL'],
    'PINK': ['PNK', 'PINK'],
    'PURPLE': ['PRP', 'PUR', 'PURPLE'],
    'YELLOW': ['YEL', 'YLW', 'YELLOW'],
    'ORANGE': ['ORG', 'ORN', 'ORANGE'],
    'TAN': ['TAN'],
    'TEAL': ['TEL', 'TEAL'],
    'CHOCOLATE': ['CHL', 'CHO', 'CHOCOLATE'],
    'COCOA': ['CCO', 'COCOA'],
    'ESPRESSO': ['ESP', 'ESPRESSO'],
    'MOCHA': ['MOC', 'MOCHA'],
    'CHARCOAL GREY': ['CHARCOAL GREY', 'CHGY'],
    'NATURAL': ['NAT', 'NATURAL'],
    'OATMEAL': ['OAT', 'OATMEAL'],
    'MUSTARD': ['MUS', 'MUSTARD'],
    'LINEN': ['LIN', 'LINEN'],
    'HONEY': ['HNY', 'HONEY'],
    'METALLIC': ['MTL', 'MET', 'METALLIC'],
    'PEWTER': ['PTR', 'PEW', 'PEWTER'],
    'BRONZE': ['BRZ', 'BRONZE'],
    'CHAMPAGNE': ['CHAMP', 'CHAMPAGNE'],
    'RUST': ['RST', 'RUST'],
    'TERRACOTTA': ['TER', 'TERRACOTTA'],
    'MULTI': ['MLT', 'MULTI'],
    'DAMASK': ['DMK', 'DAMASK'],
    'STRIPE': ['STR', 'STRIPE', 'STRIP'],
    'FLORAL': ['FLR', 'FLORAL'],
    'LEOPARD': ['LEO', 'LEOPARD'],
    'CHEETAH': ['CHT', 'CHEETAH'],
    'COWHIDE': ['COW', 'COWHIDE'],
    'BROCADE': ['BRC', 'BROCADE'],
    'JACOBEAN': ['JAC', 'JACOBEAN'],
    'MEDALLION': ['MED', 'MEDALLION'],
    'FLAMESTITCH': ['FLM', 'FLAMESTITCH'],
}


def store_label(store):
    return STORE_LABELS.get(store, store)


def store_rank(label):
    return STORE_ORDER.index(label) if label in STORE_ORDER else 99


def load_style_rows(base_dir):
    """读取款式表全量行: {SPU: [{'group','operator'}, ...]}
    同一 SPU 多行 = 该 SPU 归属多个运营组 (如 软包组/Cici + 综合组-AMZ/杨志怡),
    店铺运营规则需要按运营组分别取人, 故不能只用"一行一个SPU"的映射。"""
    import json as _json
    paths = [os.path.join(str(base_dir), '款式表_完整数据.json')]
    try:
        from process_script import BASE_DIR as _BD
        paths.append(os.path.join(str(_BD), '款式表_完整数据.json'))
    except Exception:
        pass
    for p in paths:
        if not p or not os.path.exists(p):
            continue
        try:
            data = _json.load(open(p, 'r', encoding='utf-8'))
        except Exception:
            continue
        out = defaultdict(list)
        for it in (data[1:] if isinstance(data, list) else []):
            if not isinstance(it, dict):
                continue
            spu = (it.get('spu') or '').strip()
            if not spu:
                continue
            op = (it.get('operator') or '').strip()
            if op == '/':
                op = ''
            out[spu].append({'group': (it.get('group') or '').strip(), 'operator': op})
        return out
    return {}


def make_store_op_resolver(style_rows, all_existing_map):
    """返回 resolve(spu, 店铺显示名, alt_spus=()) -> (运营组, 运营)

    Walmart / XYT / JYT : 固定人 (综合组-WLM·Chloris / 品牌组·MIA黄)
    Urban               : 取该 SPU 款式表里「综合组-AMZ」那条的运营
    DECO / HN           : 取该 SPU 款式表里「非综合组-AMZ」那条的运营 (板式组/软包组/品牌组…)
    占位运营组          : 「项目组」运营显示 '/'
    兜底①              : 该 SPU 在款式表的原值 (all_existing_map)
    兜底②              : 本体 SPU 在款式表完全无记录时, 用「上架SKU解析出的SPU」再查一遍
                          (2026-09-21 用户确认: 如 BDJS0893-KB 的 Urban 上架SKU 是 BDJH0885-KB-xxx
                           → 按 BDJH0885 的款式表记录取运营/运营组)
    """
    def _from_rows(rows, store_disp):
        if store_disp == 'Urban':
            cand = [r for r in rows if r['group'] == '综合组-AMZ'
                    and r['operator'] not in FAKE_OPERATOR]
            return ('综合组-AMZ', cand[0]['operator']) if cand else None
        if store_disp in ('DECO', 'HN'):
            pools = [
                # ① 真实运营组 (软包组/板式组/品牌组…)
                [r for r in rows if r['group'] != '综合组-AMZ'
                 and r['group'] not in FAKE_GROUP and r['operator'] not in FAKE_OPERATOR],
                # ② 放宽: 非综合组-AMZ 且有运营人
                [r for r in rows if r['group'] != '综合组-AMZ'
                 and r['operator'] not in FAKE_OPERATOR],
                # ③ 只剩「项目组」这类占位运营组: 运营组照显示, 运营用 '/' 占位
                #    (不上架/已下架/套组不上 不算候选, 交给 all_existing_map 兜底)
                [r for r in rows if r['group'] == '项目组'],
            ]
            for pool in pools:
                if pool:
                    g, o = pool[0]['group'], pool[0]['operator']
                    return (g, o or ('/' if g == '项目组' else ''))
            return None
        return None

    def _from_map(info):
        if not info:
            return None
        g = (info.get('group') or '').strip()
        o = (info.get('operator') or '').strip()
        if not g and not o:
            return None
        return (g, o or ('/' if g == '项目组' else ''))

    def resolve(spu, store_disp, alt_spus=()):
        fixed = STORE_FIXED_OP.get(store_disp)
        if fixed:
            return fixed
        hit = _from_rows(style_rows.get(spu) or [], store_disp)
        if hit:
            return hit
        hit = _from_map(all_existing_map.get(spu))
        if hit:
            return hit
        # 本体 SPU 款式表无记录 → 用「上架SKU解析出的SPU」再查
        for cand in alt_spus:
            if not cand or cand == spu:
                continue
            hit = _from_rows(style_rows.get(cand) or [], store_disp)
            if hit:
                return hit
            hit = _from_map(all_existing_map.get(cand))
            if hit:
                return hit
        return ('', '')

    return resolve


# ---- 颜色档案 (v4 archive-first) ----

def _read_local_color_dict(path):
    """读取本机 output/color_dict.json 镜像 (可能比云端更新)"""
    try:
        import json as _json
        with open(path, 'r', encoding='utf-8') as f:
            return _json.load(f)
    except Exception:
        return None


def _merge_dict_source(archive, data):
    if not isinstance(data, dict):
        return archive
    for k, v in (data.get('colors') or {}).items():
        if k and isinstance(v, dict):
            archive[str(k).strip().upper()] = {
                'name': str(v.get('name', '')).strip(),
                'status': v.get('status', '颜色'),
            }
    return archive


def load_color_archive(base_dir):
    """颜色档案 = 内置种子 + 云端 color_dict.json + 本机镜像 (人工确认, 覆盖种子)。

    合并顺序: 种子 → 云端 → 本机磁盘镜像 (后写覆盖)。
    本机镜像必须参与合并: Supabase 免费层「写入后 5-20s 读不到」,
    若只读云端, 用户在网页端确认颜色后立刻点「处理并生成报告」会读到旧档案,
    表现为「确认了却没生效」(2026-09-21 实测)。

    返回 {色码大写: {'name': 规范色名, 'status': '颜色'|'非颜色'}}
    """
    archive = {}
    try:
        from color_dict_seed import SEED_COLORS
        for k, v in SEED_COLORS.items():
            archive[str(k).strip().upper()] = {
                'name': str(v.get('name', '')).strip(),
                'status': v.get('status', '颜色'),
            }
    except Exception as e:
        print('[new_color_module] color seed load failed:', e)
    local_path = os.path.join(str(base_dir), 'output', 'color_dict.json')
    try:
        import cloud_store
        _merge_dict_source(archive, cloud_store.read_json(
            cloud_store.BUCKET_OUTPUTS, 'color_dict.json', local_path))
    except Exception as e:
        print('[new_color_module] color archive cloud load failed (fallback local/seed):', e)
    _merge_dict_source(archive, _read_local_color_dict(local_path))
    return archive


def load_transfer_groups(base_dir):
    """读取 转单*.xlsx (全部 sheet) —— 每行 = 同一实物 SKU 在不同工厂/渠道的不同名字(转单关系)。

    文件结构 (表头: 原商家SKU | 第一次转单 | 第二次转单 | 第三次转单 | 第四次转单 | 流通sku):
      同一行内的所有非空名字 = 同一个品(同颜色), 只是随工厂/渠道换名。
      行的第1列通常是原始名, 末列「流通sku」通常是当前在用名 —— 但语义上同行全部等价, 不区分角色。
      两个 sheet (增长/自营) 都读, 合并到同一套分组。

    返回 (sku_to_group: {SKU名: 组ID}, group_names: {组ID: set(SKU名)}) ; 无文件时返回 None。
    支持跨行传递合并: 若某名字同时出现在多行, 把这些行的组并成一组。
    """
    files = sorted(glob.glob(os.path.join(str(base_dir), '转单*.xlsx')))
    if not files:
        return None
    import openpyxl  # noqa: E402
    sku_to_group = {}
    group_names = {}
    next_gid = [0]

    def new_group(names):
        g = next_gid[0]
        next_gid[0] += 1
        group_names[g] = set(names)
        for n in names:
            sku_to_group[n] = g
        return g

    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=2, values_only=True):
                    names = []
                    for v in row:
                        if v is not None:
                            s = str(v).strip()
                            if s:
                                names.append(s)
                    if not names:
                        continue
                    existing = {sku_to_group[n] for n in names if n in sku_to_group}
                    if not existing:
                        new_group(names)
                    elif len(existing) == 1:
                        g = next(iter(existing))
                        for n in names:
                            if n not in sku_to_group:
                                sku_to_group[n] = g
                                group_names[g].add(n)
                    else:
                        # 跨行传递合并: 把多个已存在组合并为一个
                        g = min(existing)
                        for o in (existing - {g}):
                            for m in group_names.get(o, ()):
                                sku_to_group[m] = g
                                group_names[g].add(m)
                            group_names.pop(o, None)
                        for n in names:
                            if n not in sku_to_group:
                                sku_to_group[n] = g
                                group_names[g].add(n)
        finally:
            wb.close()

    if not sku_to_group:
        return None
    print('[new_color_module] 转单关系载入: 文件 %d 个, 分组 %d 组, 名称 %d 个'
          % (len(files), len(group_names), len(sku_to_group)))
    return sku_to_group, group_names


def _build_variant_index(archive):
    """构建「缩写/全称 → 规范色名」索引 (仅供档案未命中的写法归一)
    - 只归一为档案里已存在的规范色名, 不发明新色名
    - 歧义写法(如 CML 同时可指 CAMEL/CARAMEL)直接丢弃, 交回「待确认」
    返回: (规范色名集合, {写法: 规范色名})"""
    canon = {rec['name'].strip().upper() for rec in archive.values()
             if rec.get('status') == '颜色' and rec.get('name')}
    vmap = {}
    dup = set()
    for name, variants in STANDARD_COLOR_VARIANTS.items():
        if name not in canon:
            continue
        for v in variants:
            v = v.strip().upper()
            if v in vmap and vmap[v] != name:
                dup.add(v)
            vmap[v] = name
    for d in dup:
        vmap.pop(d, None)
    for c in canon:
        vmap.setdefault(c, c)   # 档案规范色名自身; 已被同义写法占用时不覆盖
    return canon, vmap


def _glob1(folder, pattern):
    hits = sorted(glob.glob(os.path.join(folder, pattern)))
    return hits[0] if hits else None


def _read_csv_any(path, **kw):
    """读 CSV, 自动兼容 utf-8-sig / utf-8 / gbk 编码

    上传的对照源 CSV 导出编码不固定 (Walmart/套件成员有时为 GBK),
    硬编码 encoding='utf-8-sig' 会抛 UnicodeDecodeError 导致整个新色判定失败。
    """
    last = None
    for enc in ('utf-8-sig', 'utf-8', 'gbk'):
        try:
            return pd.read_csv(path, encoding=enc, dtype=str, **kw).fillna('')
        except UnicodeDecodeError as e:
            last = e
    raise last


def _is_nondeleted(store, status):
    """非删除 = 已上架链接存在 (有资格作为母体, 且计入基准)"""
    if not status:
        return False
    if store == WALMART_STORE[0]:
        return status in ('PUBLISHED', 'UNPUBLISHED')
    return status in ('在售', '停售')


def compute_new_colors(base_dir, color_archive=None):
    """返回 dict:
       spu_rows:  [ {SPU,店铺,运营组,运营,已上架店铺,新色数,新色列表,待上架SKU数}, ... ]
                  一行 = 一个 SPU × 一个店铺 (同SPU多店铺有上架 → 多条记录, 2026-09-21 用户确认)
                  运营组/运营 按店铺规则取值 (Walmart/XYT/JYT 固定人; Urban 取综合组-AMZ; DECO/HN 取非综合组-AMZ)
       sku_rows:  [ {SPU,上架SKU,库存SKU,颜色,店铺,上架状态(已上架/未上架/非颜色),颜色判定,
                     在库库存,在途库存,在途最早预计到仓,库存状态,数量,预计到仓}, ... ]
                  展开用。已上架=「上架SKU+库存SKU」一对一行; 未上架(新色)=仅库存SKU(上架SKU为空);
                  非颜色=色码被档案判定为非颜色(不参与新色判定)
                  在库库存/在途库存 分开列出; 预计到仓 = 在途最早预计到仓时间
       new_colors: 仅未上架(待上架新色)行, 供导出/核对 (含 颜色判定 字段)
       pending_colors: [ {颜色码,SPU数,SKU数,新色行数,示例SKU,示例SPU,涉及店铺}, ... ]
                  档案中不存在、需人工确认的新颜色码 (网页端「颜色待确认」面板)
       warn_unmatched: [ (店铺, 未命中套件的非删除平台SKU, 状态), ... ]
       warn_unparse: [ (店铺, SKU, SPU, 状态), ... ]  (已上架SPU的在库/在途SKU无法解析颜色)
    """
    archive = color_archive if color_archive is not None else load_color_archive(base_dir)
    _canon_set, variant_map = _build_variant_index(archive)
    auto_mapped = defaultdict(set)   # 档案未命中但通过缩写/全称归一识别的原始码 -> 规范色名

    def map_color(raw):
        """原始色码 → (规范色名, 判定); 判定: '颜色'/'非颜色'/'待确认'/''"""
        if not raw:
            return '', ''
        t = str(raw).strip().upper()
        # 内置非颜色(配件)色码: 不进颜色档案也能直接判非颜色, 避免被当成待确认新色
        if t in _BUILTIN_NON_COLORS:
            return '', '非颜色'
        rec = archive.get(t)
        if rec:
            if rec.get('status') == '非颜色' or not rec.get('name'):
                return '', '非颜色'
            name = rec['name'].strip().upper()
            # 规范色名本身也过一遍归一 (档案里可能写成别名)
            return (variant_map.get(name, name), '颜色')
        # 档案未命中 → 缩写/全称归一 (BEM0908-brown ≡ BEM0908-BRN)
        c = variant_map.get(t)
        if c:
            # 仅当真正发生了写法变化才记入溯源 (自我映射如 CREAM→CREAM 不是"归一", 记录只会造成噪音)
            if c != t:
                auto_mapped[t].add(c)
            return (c, '颜色')
        return (t, '待确认')
    # ---- 定位文件 ----
    f_listing = _glob1(base_dir, '*Listing*.xlsx')
    f_wm = _glob1(base_dir, '*ItemReport*.csv')
    f_suite = _glob1(base_dir, '*套件成员*.csv')
    f_inv = _glob1(base_dir, '库存清单*.xlsx')
    f_transit = _glob1(base_dir, '动态库存表*.xlsx')
    missing = [n for n, p in [('Listing', f_listing), ('ItemReport', f_wm),
                              ('套件成员', f_suite), ('库存清单', f_inv),
                              ('动态库存表', f_transit)] if not p]
    if missing:
        raise FileNotFoundError('缺少数据文件: ' + ', '.join(missing))

    # ---- 转单关系 (可选数据源) ----
    # 转单*.xlsx: 同一行的所有名字 = 同一实物 SKU (不同工厂/渠道换名), 往往落到不同 SPU。
    # 语义: 同组内任一名字在某店铺非删除非淘汰上架 → 全组该颜色都视为已上架 → 不算新色。
    _transfer = load_transfer_groups(base_dir)
    transfer_sku_to_group = (_transfer[0] if _transfer else {}) or {}
    _transfer_enabled = bool(transfer_sku_to_group)

    from process_script import (  # noqa: E402
        extract_spu_from_sku, parse_sku_color_suffix, load_existing_spu_list,
    )

    style = load_existing_spu_list()
    known_spus = set(style['spu_set']) | set(style['spu_to_aliases'].keys()) | set(style['alias_to_spu'].keys())
    alias_to_spu = dict(style['alias_to_spu'])

    def norm_spu(spu):
        # 子件规则: 若该码本身就是款式表 A 列的独立 SPU (如套件 Z6SFRCCT0685 的子件
        # Z6SFXW0526 自己也是 SPU), 不再强行归并进套件 SPU —— 组合件与单件各自独立
        # 判新色 (不同 SPU 各自出待上架新色); 仅「非独立SPU的别名码」才归并到母体。
        if spu in style['spu_set']:
            return spu
        return alias_to_spu.get(spu, spu)

    def sku_to_spu_color(sku):
        spu = extract_spu_from_sku(sku, known_spus)
        color, _sfx = parse_sku_color_suffix(sku, spu)
        return norm_spu(spu), (color.upper() if color else '')

    # ---- 套件成员映射 (含生命周期) ----
    suite = _read_csv_any(f_suite)
    suite = suite[suite['成员货品'].str.strip() != '']
    suite_dedup = suite.drop_duplicates(subset=['店铺', '14K货号', '成员货品'])
    suite_members = defaultdict(set)
    suite_lc = {}                       # (store, 14K货号/平台SKU) -> 美国生命周期状态
    member_to_k14 = defaultdict(list)   # (store, 成员货品) -> [14K货号, ...] (成员可能跨多个货号)
    for _, r in suite_dedup.iterrows():
        st = r['店铺'].strip()
        k14 = r['14K货号'].strip()
        mem = r['成员货品'].strip()
        lc = str(r.get('美国生命周期状态', '') or '').strip()
        suite_members[(st, k14)].add(mem)
        if lc:
            suite_lc[(st, k14)] = lc
        member_to_k14[(st, mem)].append(k14)
    # ---- Listing 全量状态 (保留状态用于母体判定) ----
    listing_status = {}  # (store, platform_sku) -> status
    import openpyxl  # noqa: E402
    wb = openpyxl.load_workbook(f_listing, read_only=True)
    for row in wb.active.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        msku, status, shop = str(row[0]).strip(), str(row[1] or '').strip(), str(row[2] or '').strip()
        if shop not in AMAZON_STORE_MAP:
            continue
        store = AMAZON_STORE_MAP[shop][0]
        listing_status[(store, msku)] = status
    wb.close()
    wm = _read_csv_any(f_wm)
    for _, r in wm.iterrows():
        sku = str(r['SKU'] or '').strip()
        if not sku:
            continue
        st = str(r['Publish Status'] or '').strip()
        listing_status[(WALMART_STORE[0], sku)] = st

    # ---- 在库 + 在途 池 (提前读取, 用于生命周期解析 + 后续建池) ----
    inv = pd.read_excel(f_inv, engine='openpyxl')
    inv = inv[pd.notna(inv['SKU'])]
    tr = pd.read_excel(f_transit, engine='openpyxl')
    tr = tr[pd.notna(tr['SKU'])]
    tr['SKU'] = tr['SKU'].astype(str).str.strip()

    def _find_lifecycle_col(df):
        for c in df.columns:
            if '生命周期' in str(c):
                return c
        return None

    # 库存清单 生命周期状态 (在库, 有列才用)
    inv_lc = {}
    _lc_col = _find_lifecycle_col(inv)
    if _lc_col:
        for _, r in inv.iterrows():
            sku = str(r['SKU'] or '').strip()
            lc = str(r.get(_lc_col, '') or '').strip()
            if sku and lc:
                inv_lc[sku] = lc
    # 动态库存表 生命周期状态 (在途)
    transit_lc = {}
    _lc_col2 = _find_lifecycle_col(tr)
    if _lc_col2:
        for _, r in tr.iterrows():
            sku = str(r['SKU'] or '').strip()
            lc = str(r.get(_lc_col2, '') or '').strip()
            if sku and lc:
                transit_lc[sku] = lc

    def is_eliminated(lifecycle):
        """淘汰判定: 生命周期含"淘汰" 或 等于"下市期" → 已淘汰 (死货/小黑屋/空值 不淘汰)"""
        if not lifecycle:
            return False
        v = str(lifecycle).strip()
        return ('淘汰' in v) or (v == '下市期')

    def resolve_lifecycle(store, sku):
        """SKU 生命周期解析 (优先级: 动态库存表[在途] → 库存清单[在库] → 套件表[美国生命周期状态])
        返回生命周期字符串 (空=未查到=不淘汰)"""
        if sku in transit_lc:
            return transit_lc[sku]
        if sku in inv_lc:
            return inv_lc[sku]
        lc = suite_lc.get((store, sku))
        if lc:
            return lc
        for k14 in member_to_k14.get((store, sku), []):
            lc2 = suite_lc.get((store, k14))
            if lc2:
                return lc2
        return ''

    def kit_all_nondeleted(store, sku):
        """规则2: 若 库存SKU 是某上架SKU(平台SKU)的套件成员, 其所属套件组(全部库存子SKU)
        必须全非淘汰, 才算需要上架的新色"""
        k14s = member_platforms.get((store, sku), set())
        for k14 in k14s:
            for m in suite_members.get((store, k14), set()):
                if is_eliminated(resolve_lifecycle(store, m)):
                    return False
        return True

    # ---- 分店铺 基准颜色 + 母体状态 ----
    baseline_colors = defaultdict(lambda: defaultdict(set))   # store -> spu -> {colors of 非删除且非淘汰}
    all_statuses = defaultdict(lambda: defaultdict(set))      # store -> spu -> {status}
    member_platforms = defaultdict(set)  # (store, 子件/库存SKU) -> {该子件对应的已上架平台SKU}
    platform_status = {}                 # (store, 平台SKU/14K货号) -> 原始后台状态 (供前端"上架SKU汇总"展示)
    group_listed_names = defaultdict(lambda: defaultdict(set))  # store -> 转单组ID -> 本店已上架的名字
    warn_unmatched = []
    for (store, k14), status in listing_status.items():
        platform_status[(store, k14)] = status
        plat_spu = sku_to_spu_color(k14)[0]
        members = suite_members.get((store, k14))
        if members:
            # 组合件自身也算一个"成员": 套件(组合件)自己的上架颜色要计入其 SPU 基准。
            # 否则组合件已上架、子件单独也上架时, 组合件自身的颜色不进基准 → 误报新色。
            members = set(members) | {k14}
        else:
            if _is_nondeleted(store, status) and not is_eliminated(resolve_lifecycle(store, k14)):
                warn_unmatched.append((store_label(store), k14, status))
            members = {k14}  # 回退: 平台SKU自身按库存级SKU处理
        for m in members:
            spu, color = sku_to_spu_color(m)
            if not spu:
                continue
            all_statuses[store][spu].add(status)
            # 规则4(母体资格): 仅计入「非删除 且 非淘汰」的已上架颜色作为基准;
            # 仅已删除 / 仅淘汰 / 无有效在售 → 整行无新色 (下方判定循环用 bl 是否为空拦)
            # 仅「同SPU」成员才计入基准: 跨SPU成员(如套件子件属另一SPU)若计入会污染他SPU基准,
            # 导致规则2(套件新色)无法触发 —— 跨SPU成员的"是否上架"交由该SPU自身 Listing 决定。
            if _is_nondeleted(store, status) and not is_eliminated(resolve_lifecycle(store, k14)):
                raw_c = color          # 原始色码 (map_color 归一前)
                color, _cstat = map_color(color)
                member_platforms[(store, m)].add(k14)
                if color and spu == plat_spu:
                    baseline_colors[store][spu].add(color)
                # 转单关系: 把该已上架名字登记到所属转单组。
                # ⚠ 不比对色名: 转单组=同一实物, 换厂/换渠道时连色码写法都会变
                # (如 NSMK1427-WAL[WALNUT] ↔ NSQY1845-AUB[AUBURN] 是同一件货),
                # 同组任一名字在本店已上架(非删除非淘汰) → 全组都算已上架。
                # 平台SKU(k14) 与库存子件(m) 两侧都登记: 转单表里两类名字都可能出现。
                if _transfer_enabled:
                    for _nm in (k14, m):
                        _g = transfer_sku_to_group.get(_nm)
                        if _g is not None:
                            group_listed_names[store][_g].add(_nm)

    # ---- 在库 + 在途 池 (inv/tr 已在上文提前读取并提取生命周期) ----
    inv_qty = inv.groupby(inv['SKU'].astype(str).str.strip())['库存量'].sum().to_dict()

    def fmt_eta(v):
        if pd.isna(v):
            return ''
        if hasattr(v, 'strftime'):
            return v.strftime('%Y-%m-%d')
        return str(v).strip()[:10]

    tr_g = tr.groupby('SKU').agg(qty=('数量', 'sum'), eta=('预计到仓时间', 'min')).reset_index()
    transit_qty = dict(zip(tr_g['SKU'], tr_g['qty']))
    transit_eta = {r['SKU']: fmt_eta(r['eta']) for _, r in tr_g.iterrows()}

    pool = {}
    for sku, q in inv_qty.items():
        pool[sku] = [0.0, 0.0, '']            # [在库量, 在途量, 在途最早ETA]
    for sku, q in transit_qty.items():
        if sku not in pool:
            pool[sku] = [0.0, 0.0, '']
        pool[sku][1] += float(q or 0)
        pool[sku][2] = transit_eta.get(sku, '')
    for sku, v in pool.items():
        v[0] = float(inv_qty.get(sku, 0) or 0)

    def _qty(v):
        f = float(v)
        return int(f) if abs(f - int(f)) < 1e-9 else round(f, 2)

    def pool_stat(v):
        inv_q, tr_q = v[0], v[1]
        if inv_q and tr_q:
            return '在库+在途'
        if tr_q:
            return '在途'
        return '在库'

    pool_parsed = {}
    for sku, v in pool.items():
        spu, raw_color = sku_to_spu_color(sku)
        color, cstat = map_color(raw_color)
        inv_q, tr_q, eta = _qty(v[0]), _qty(v[1]), v[2]
        pool_parsed[sku] = (spu, color, cstat, raw_color, pool_stat(v), inv_q, tr_q, eta)

    # SPU -> [在库/在途 SKU]
    spu_skus = defaultdict(list)
    for sku, pr in pool_parsed.items():
        if pr[0]:
            spu_skus[pr[0]].append(sku)

    # ---- 判定: SKU×店铺 明细(已上架/未上架/非颜色) + 新色 + 待确认颜色 ----
    sku_rows = []
    new_color_rows = []
    warn_unparse = []
    pending = defaultdict(lambda: {'spus': set(), 'skus': set(), 'stores': set(), 'rows': 0})
    for store in sorted(all_statuses):
        st_map = all_statuses[store]
        bl_map = baseline_colors.get(store, {})
        for spu, statuses in st_map.items():
            bl = bl_map.get(spu, set())
            if not bl:
                continue  # 情形② / 仅淘汰 / 无有效在售 → 该店铺该SPU当前无需上架, 跳过(规则4)
            for sku in spu_skus.get(spu, []):
                _spu, color, cstat, raw_color, st, inv_q, tr_q, eta = pool_parsed[sku]
                if not color and cstat != '非颜色':
                    warn_unparse.append((store, sku, spu, st))
                    continue
                listed_by_self = color in bl
                # 转单关系: 本库存SKU 若属某转单组, 且同组任一名字在本店已上架 → 视为已上架(不算新色)。
                # 不比对色名(转单组=同一实物, 色码写法可随换厂变化, 如 WAL↔AUBURN)。
                _g = transfer_sku_to_group.get(sku) if _transfer_enabled else None
                covered_transfer = bool(
                    _g is not None and group_listed_names.get(store, {}).get(_g)
                )
                listed = listed_by_self or covered_transfer
                plats = sorted(member_platforms.get((store, sku), set()))
                store_disp = store_label(store)

                def _mk(status, plat, color_val, backend='', transfer=False):
                    return {
                        'SPU': spu,
                        '上架SKU': plat,
                        '库存SKU': sku,
                        '颜色': color_val,
                        '店铺': store_disp,
                        '上架状态': status,
                        '颜色判定': cstat,
                        '后台状态': backend,
                        '转单覆盖': ('是' if transfer else ''),
                        '在库库存': inv_q,
                        '在途库存': tr_q,
                        '在途最早预计到仓': eta,
                        '库存状态': st,
                        '数量': (inv_q or 0) + (tr_q or 0),
                        '预计到仓': eta,
                    }

                # 规则1: 淘汰色(生命周期含"淘汰"或"下市期") → 无需安排上架, 不算新色
                if is_eliminated(resolve_lifecycle(store, sku)):
                    sku_rows.append(_mk('已淘汰', '', raw_color))
                    continue
                if cstat == '非颜色':
                    # 档案判定为非颜色: 不参与新色判定, 明细行灰色显示
                    sku_rows.append(_mk('非颜色', '', raw_color))
                    continue
                if listed and plats:
                    # 已上架: 按「上架SKU + 库存SKU」维度, 一个平台SKU对应一行
                    for p in plats:
                        sku_rows.append(_mk('已上架', p, color,
                                            platform_status.get((store, p), ''),
                                            transfer=(covered_transfer and not listed_by_self)))
                elif listed:
                    # 同色其他子件/同组转单名已上架(本件未单独挂链接) → 仍算已上架, 无对应上架SKU
                    sku_rows.append(_mk('已上架', '', color,
                                        transfer=(covered_transfer and not listed_by_self)))
                else:
                    # 规则2: 套件新色 — 上架SKU对应的所有库存子SKU 都非淘汰, 才算需要上架的新色
                    if not kit_all_nondeleted(store, sku):
                        sku_rows.append(_mk('套件淘汰', '', color))
                        continue
                    # 未上架 = 待上架新色: 按库存SKU维度, 带在库/在途/在途最早到仓
                    sku_rows.append(_mk('未上架', '', color))
                    new_color_rows.append({
                        'SPU': spu,
                        '店铺': store_disp,
                        '新色': color,
                        '颜色判定': cstat,
                        '库存SKU': sku,
                        '在库库存': inv_q,
                        '在途库存': tr_q,
                        '在途最早预计到仓': eta,
                        '库存状态': st,
                        '数量': (inv_q or 0) + (tr_q or 0),
                        '预计到仓': eta,
                    })
                    if cstat == '待确认':
                        e = pending[color]
                        e['spus'].add(spu)
                        e['skus'].add(sku)
                        e['stores'].add(store_label(store))
                        e['rows'] += 1

    pending_colors = [
        {
            '颜色码': k,
            'SPU数': len(v['spus']),
            'SKU数': len(v['skus']),
            '新色行数': v['rows'],
            '示例SKU': sorted(v['skus'])[0],
            '示例SPU': sorted(v['spus'])[0],
            '涉及店铺': ', '.join(sorted(v['stores'])),
        }
        for k, v in pending.items()
    ]
    pending_colors.sort(key=lambda x: (-x['新色行数'], x['颜色码']))

    # ---- 上架SKU 解析出的 SPU (供款式表查不到本体SPU时兜底查运营) ----
    alt_spu_map = defaultdict(list)   # (SPU, 店铺显示名) -> [上架SKU的SPU, ...]
    for r in sku_rows:
        us = r.get('上架SKU')
        if not us:
            continue
        s2 = norm_spu(extract_spu_from_sku(us, known_spus))
        if s2 and s2 != r['SPU']:
            lst = alt_spu_map[(r['SPU'], r['店铺'])]
            if s2 not in lst:
                lst.append(s2)

    # ---- SPU 主行聚合: 一行 = 一个 SPU × 一个店铺 (同SPU多店上架 → 多行) ----
    agg = defaultdict(lambda: {'新色': set(), '待上架数': 0})   # (SPU, 店铺显示名) -> 聚合
    for r in sku_rows:
        if r['上架状态'] != '未上架':
            continue
        e = agg[(r['SPU'], r['店铺'])]
        e['新色'].add(r['颜色'])
        e['待上架数'] += 1

    all_existing_map = style.get('all_existing_map') or {}
    style_rows = load_style_rows(base_dir)
    resolver = make_store_op_resolver(style_rows, all_existing_map)

    spu_rows = []
    for (spu, store_disp), e in agg.items():
        if not e['新色']:
            continue  # 无待上架新色的店铺记录不进跟进表
        info = all_existing_map.get(spu) or {}
        grp, op = resolver(spu, store_disp, alt_spu_map.get((spu, store_disp), []))
        spu_rows.append({
            'SPU': spu,
            '店铺': store_disp,
            '运营组': grp,
            '运营': op,
            '已上架店铺': store_disp,   # 一行 = SPU×店铺, 已上架店铺只显示本行店铺 (2026-09-21 用户确认)
            '新色数': len(e['新色']),
            '新色列表': ', '.join(sorted(e['新色'])),
            '待上架SKU数': e['待上架数'],
            '款式表运营组': info.get('group', ''),
            '款式表运营': info.get('operator', ''),
        })
    spu_rows.sort(key=lambda x: (-x['新色数'], x['SPU'], store_rank(x['店铺'])))

    # 运营组/运营下拉选项 (与新品跟进同源: 款式表全集)
    groups = sorted({v.get('group', '') for v in all_existing_map.values() if v.get('group')})
    operators = sorted({v.get('operator', '') for v in all_existing_map.values() if v.get('operator')})

    return {
        'spu_rows': spu_rows,
        'sku_rows': sku_rows,
        'new_colors': new_color_rows,
        'pending_colors': pending_colors,
        'auto_colors': [
            {'原始码': k, '归一为': ', '.join(sorted(v))}
            for k, v in sorted(auto_mapped.items())
        ],
        'style_options': {'groups': groups, 'operators': operators},
        'warn_unmatched': warn_unmatched,
        'warn_unparse': warn_unparse,
    }


if __name__ == '__main__':
    import json
    res = compute_new_colors(os.path.dirname(os.path.abspath(__file__)))
    print('SPU 主行(SPU×店铺):', len(res['spu_rows']), '| 涉及 SPU 数:', len({r['SPU'] for r in res['spu_rows']}))
    print('SKU 明细行:', len(res['sku_rows']), '(已上架[上架SKU+库存SKU] + 未上架[库存SKU] + 非颜色)')
    print('  其中已上架:', sum(1 for r in res['sku_rows'] if r['上架状态'] == '已上架'))
    print('  其中未上架:', sum(1 for r in res['sku_rows'] if r['上架状态'] == '未上架'))
    print('  其中非颜色:', sum(1 for r in res['sku_rows'] if r['上架状态'] == '非颜色'))
    print('待上架新色行:', len(res['new_colors']))
    print('  其中颜色待确认行:', sum(1 for r in res['new_colors'] if r.get('颜色判定') == '待确认'))
    print('待确认颜色码数:', len(res['pending_colors']))
    print('未命中套件(非删除):', len(res['warn_unmatched']))
    print('颜色无法解析:', len(res['warn_unparse']))
    print(json.dumps(res['spu_rows'][:3], ensure_ascii=False, indent=2))
